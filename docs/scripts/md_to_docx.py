"""Convierte un Markdown simple (headers, **bold**, `code`, ---, listas, [link](url)) a DOCX via python-docx.
Uso: python md_to_docx.py entrada.md salida.docx "Titulo del documento"
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

INK = RGBColor(0x1C, 0x24, 0x30)
ACCENT = RGBColor(0x12, 0x3A, 0x7A)
ACCENT_SOFT = RGBColor(0x2A, 0x5D, 0xB0)
MUTED = RGBColor(0x4C, 0x55, 0x68)

INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline_runs(paragraph, text: str):
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = ACCENT
        elif chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        elif chunk.startswith("[") and "](" in chunk:
            label = chunk[1:chunk.index("]")]
            run = paragraph.add_run(label)
            run.italic = True
            run.font.color.rgb = ACCENT_SOFT
        else:
            paragraph.add_run(chunk)


def set_cell_shading(cell, color_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): color_hex})
    tcPr.append(shd)


def build(md_text: str, title: str) -> Document:
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK

    lines = md_text.replace("\r\n", "\n").split("\n")
    in_ul = False

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped == "":
            in_ul = False
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pPr = p._p.get_or_add_pPr()
            pBdr = pPr.makeelement(qn("w:pBdr"), {})
            bottom = pPr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "C9C2B3"})
            pBdr.append(bottom)
            pPr.append(pBdr)
            in_ul = False
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            in_ul = False
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                h = doc.add_heading(level=1)
                run = h.add_run(text)
                run.font.color.rgb = ACCENT
                run.font.size = Pt(20)
            elif level == 2:
                h = doc.add_heading(level=2)
                run = h.add_run(text)
                run.font.color.rgb = ACCENT
                run.font.size = Pt(15)
            else:
                h = doc.add_heading(level=3)
                run = h.add_run(text)
                run.font.color.rgb = ACCENT_SOFT
                run.font.size = Pt(12)
            continue

        m = re.match(r"^\|(.+)\|$", stripped)
        if m:
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, m.group(1))
            in_ul = True
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline_runs(p, stripped)
        in_ul = False

    return doc


def main():
    if len(sys.argv) < 4:
        print('Uso: python md_to_docx.py entrada.md salida.docx "Titulo"')
        sys.exit(1)

    src, dst, title = sys.argv[1], sys.argv[2], sys.argv[3]

    with open(src, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = build(md_text, title)
    doc.core_properties.title = title
    doc.save(dst)

    print(f"DOCX generado: {dst}")


if __name__ == "__main__":
    main()
