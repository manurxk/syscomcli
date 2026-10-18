"""Genera la plantilla de ejemplo del manual de usuario (IEEE Std 1063-2001 / ISO/IEC 26514:2008).

Muestra cada elemento del estándar con un capítulo real de la guía de uso (Notas de crédito) y sus capturas.
El formato (estilos, portada, control, índice, encabezado y pie) está en app/services/manual_word_service.py,
el mismo que usa el sistema para los manuales que se descargan.

Uso: python docs/scripts/generar_plantilla_manual.py
Salida: docs/estandar/plantilla_manual.docx (se sobrescribe).

El índice y la numeración de figuras son campos de Word: al abrir el archivo, seleccionar todo (Ctrl+A)
y presionar F9 para actualizarlos.
"""
import os
import sys

from docx.shared import Cm

sys.path.insert(0, os.path.dirname(__file__))
from generar_manual_seguridad_docx import RAIZ, LOGO, cargar_servicio  # noqa: E402

mw = cargar_servicio()

SALIDA = os.path.join(RAIZ, 'docs', 'estandar', 'plantilla_manual.docx')
CAPTURAS = os.path.join(RAIZ, 'app', 'static', 'img', 'ayuda', 'notas_credito')
TITULO_DOC = 'Guía de uso del sistema'
FECHA = '17/10/2026'


def _figura(doc, archivo, descripcion):
    """Captura centrada con pie «Figura N.M.» (N = capítulo, M = número de figura en el capítulo)."""
    ruta = os.path.join(CAPTURAS, archivo)
    if os.path.isfile(ruta):
        doc.add_paragraph('', style='Figura').add_run().add_picture(ruta, width=Cm(15))
    pie = doc.add_paragraph('', style='Caption')
    r = pie.add_run('Figura ')
    r.bold = True
    r.italic = False
    mw.campo(pie, 'STYLEREF 1 \\s', '1')
    pie.add_run('.').bold = True
    mw.campo(pie, 'SEQ Figura \\* ARABIC \\s 1', '1')
    pie.add_run('. ').bold = True
    pie.add_run(descripcion)


def _pasos(doc, pasos):
    for paso in pasos:
        p = doc.add_paragraph('', style='List Number')
        for parte, negrita in paso:
            p.add_run(parte).bold = negrita


def _contenido(doc):
    # 1. Introducción (IEEE 1063: objetivo, destinatarios, cómo usar el documento, convenciones)
    doc.add_heading('Introducción', level=1)
    doc.add_heading('Objetivo', level=2)
    doc.add_paragraph('Este manual explica cómo usar el sistema SYSCLIN para las tareas del perfil Ventas: '
                      'caja, facturación, notas de crédito y débito, cuentas a cobrar y cobranzas.')
    doc.add_heading('Destinatarios', level=2)
    doc.add_paragraph('Personal de caja y facturación de la clínica. No se necesitan conocimientos técnicos; '
                      'alcanza con saber usar un navegador web.')
    doc.add_heading('Convenciones del documento', level=2)
    mw.tabla(doc, ['Convención', 'Significado'], [
        [('Negrita',), 'Botón, campo o menú del sistema que se debe presionar o completar.'],
        ['Pasos numerados', 'Procedimiento que se hace en ese orden.'],
        ['Nota', 'Información que conviene tener en cuenta.'],
        ['Advertencia', 'Acción que no se puede deshacer o que puede causar un error.'],
    ], [4.5, 12])

    # 2. Capítulo modelo (contenido real de la guía de uso)
    doc.add_heading('Notas de crédito', level=1)
    doc.add_paragraph('La nota de crédito corrige a la baja una factura ya emitida: una devolución, un descuento '
                      'que no se aplicó o un cobro de más. Se anota sola en el Libro de Ventas.')
    doc.add_heading('Emitir una nota de crédito', level=2)
    _figura(doc, '01_listado.jpg', 'Pantalla Notas de Crédito, botón Nueva Nota de Crédito')
    _pasos(doc, [
        [('Presione ', False), ('Nueva Nota de Crédito', True), ('.', False)],
        [('Elija la ', False), ('Factura', True), (' que va a corregir y presione ', False),
         ('Seleccionar y continuar', True), ('.', False)],
        [('Escriba el ', False), ('Motivo', True), (' (es obligatorio).', False)],
        [('Presione ', False), ('Agregar ítem', True),
         (' por cada concepto que se descuenta: descripción, cantidad y precio unitario.', False)],
        [('Presione ', False), ('Guardar Nota de Crédito', True), ('.', False)],
    ])
    _figura(doc, '02_nueva.jpg', 'Formulario Nueva Nota de Crédito')
    doc.add_heading('Efecto sobre la deuda del paciente', level=2)
    doc.add_paragraph('Si la factura tiene saldo pendiente de pago, la nota de crédito lo descuenta sola en su '
                      'Cuenta a Cobrar. Si con eso la deuda llega a cero, la cuenta queda como cobrada.')
    mw.nota(doc, 'Nota', 'La nota de crédito no puede ser mayor a lo que falta pagar de la factura. Si el paciente '
                         'ya pagó y hay que devolverle dinero, esa devolución se hace fuera del sistema.')
    doc.add_heading('Anular una nota de crédito', level=2)
    doc.add_paragraph('En la lista, abra el detalle de la nota y presione Anular. El monto que había descontado '
                      'vuelve a sumarse a la deuda.')
    mw.nota(doc, 'Advertencia', 'Al anular la nota, el monto que había descontado vuelve a la deuda del paciente. '
                                'Verifique la nota antes de anularla.')

    # 3. Mensajes del sistema (IEEE 1063: mensajes de error y solución de problemas)
    doc.add_heading('Mensajes del sistema', level=1)
    mw.tabla(doc, ['Mensaje', 'Causa', 'Qué hacer'], [
        ['La nota de crédito debe tener al menos un ítem.', 'Se presionó Guardar sin agregar ítems.',
         'Agregar al menos un ítem con Agregar ítem.'],
        ['Cada ítem debe tener una descripción.', 'Un ítem quedó con la descripción vacía.',
         'Completar la descripción del ítem.'],
        ['La nota ya está anulada o no se pudo anular.', 'La nota se anuló antes.',
         'Actualizar la lista y revisar el estado de la nota.'],
    ], [6, 5, 5.5])

    # 4. Glosario
    doc.add_heading('Glosario', level=1)
    mw.tabla(doc, ['Término', 'Definición'], [
        [('KuDE',), 'Representación impresa del documento electrónico (SIFEN).'],
        [('SIFEN',), 'Sistema Integrado de Facturación Electrónica Nacional (SET).'],
        [('Timbrado',), 'Autorización de la SET para emitir comprobantes en un período.'],
    ], [4, 12.5])

    # Anexo: ejemplo de bloque de código (solo para manuales técnicos)
    doc.add_paragraph('Anexo A. Ejemplo de código', style='Título de anexo')
    doc.add_paragraph('Los comandos y scripts se escriben con el estilo Código, en Consolas 9:')
    for linea in ('docker compose -f dock/docker-compose.yml ps', 'docker restart syscomcli_app'):
        doc.add_paragraph(linea, style='Código')


def main():
    doc = mw.nuevo_documento()
    mw.encabezado_y_pie(doc.sections[0], TITULO_DOC)
    mw.portada(doc, LOGO, TITULO_DOC, 'Manual de usuario · Perfil Ventas',
               ('Sistema SYSCLIN de gestión clínica', f'Versión {mw.VERSION}', FECHA))
    mw.control(doc, [
        ('Documento', f'{TITULO_DOC} (manual de usuario)'),
        ('Perfil', 'Ventas'),
        ('Versión', mw.VERSION),
        ('Fecha de generación', FECHA),
        ('Referencia', 'IEEE Std 1063-2001, Software User Documentation'),
    ], FECHA)
    mw.indice(doc, con_figuras=True)
    _contenido(doc)
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    with open(SALIDA, 'wb') as f:
        f.write(mw.guardar(doc, TITULO_DOC))
    print(f'Plantilla generada: {SALIDA}')


if __name__ == '__main__':
    main()
